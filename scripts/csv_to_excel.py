#!/usr/bin/env python3
"""
csv_to_excel.py - convert the C program's two CSV outputs into a single
Excel workbook, one sheet per instructor, laid out like the sample
CLASS SCHEDULE / OFFICE HOURS table from the project handout.

Usage: csv_to_excel.py <class_schedule.csv> <office_hours.csv> <output.xlsx>

This is invoked from the Makefile's `run` target; it is NOT required to
build the C program itself (that stands entirely on its own -- the .xlsx is
purely a presentation step over files the C program already wrote).
Requires the `openpyxl` package: pip install openpyxl
"""
import csv
import sys
from collections import defaultdict

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.stderr.write(
        "error: the 'openpyxl' package is required for this script.\n"
        "        install it with:  pip install openpyxl\n"
    )
    sys.exit(1)

DAY_ORDER = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]

TITLE_FONT = Font(bold=True, size=14)
SECTION_FONT = Font(bold=True, size=12, color="FFFFFF")
SECTION_FILL = PatternFill("solid", fgColor="2F5496")
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
NOTE_FONT = Font(italic=True, size=9)


def sanitize_sheet_name(name: str) -> str:
    for ch in '[]:*?/\\':
        name = name.replace(ch, '')
    return name[:31] or "Sheet"


def read_csv_grouped_by_instructor(path):
    grouped = defaultdict(list)
    with open(path, newline='', encoding='utf-8') as f:
        for row in csv.DictReader(f):
            grouped[row['Instructor']].append(row)
    return grouped


def section_header(ws, text, width):
    """Appends a merged, filled section-title row and returns its row number."""
    ws.append([text] + [None] * (width - 1))
    row = ws.max_row
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
    cell = ws.cell(row=row, column=1)
    cell.font = SECTION_FONT
    cell.fill = SECTION_FILL
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 18
    return row


def table_header(ws, labels):
    """Appends a bold/filled table-header row and returns its row number."""
    ws.append(labels)
    row = ws.max_row
    for col in range(1, len(labels) + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
    return row


def build_sheet(wb, instructor, courses, office_blocks):
    ws = wb.create_sheet(sanitize_sheet_name(instructor))

    ws.append([instructor])
    ws.cell(row=ws.max_row, column=1).font = TITLE_FONT
    ws.append([])  # blank spacer row

    section_header(ws, "CLASS SCHEDULE", 5)
    table_header(ws, ["Code", "Title", "Days", "Times", "Room"])

    courses.sort(key=lambda c: c['Start'])
    for c in courses:
        ws.append([c['Code'], c['Title'], c['Days'],
                   f"{c['Start']} - {c['End']}", c['Room']])

    ws.append([])  # blank spacer row
    section_header(ws, "OFFICE HOURS", 5)

    ws.append(["NOTE: Instructors are not to be interrupted during class times."])
    note_row = ws.max_row
    ws.cell(row=note_row, column=1).font = NOTE_FONT
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=5)

    table_header(ws, ["Day", "Start", "End", "Related Course", ""])

    office_blocks.sort(key=lambda b: (DAY_ORDER.index(b['Day']) if b['Day'] in DAY_ORDER else 99,
                                       b['Start']))
    for b in office_blocks:
        ws.append([b['Day'], b['Start'], b['End'], b['RelatedCourse'] or "General", ""])

    widths = (14, 34, 8, 20, 10)
    for col_idx, width in zip(range(1, 6), widths):
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def main() -> int:
    if len(sys.argv) != 4:
        sys.stderr.write(f"Usage: {sys.argv[0]} <class_schedule.csv> <office_hours.csv> <output.xlsx>\n")
        return 1

    courses_path, office_path, out_path = sys.argv[1], sys.argv[2], sys.argv[3]

    courses_by_instructor = read_csv_grouped_by_instructor(courses_path)
    office_by_instructor = read_csv_grouped_by_instructor(office_path)

    instructors = sorted(set(courses_by_instructor) | set(office_by_instructor))
    if not instructors:
        sys.stderr.write(f"warning: no rows found in '{courses_path}' or '{office_path}'\n")

    wb = Workbook()
    wb.remove(wb.active)

    for instructor in instructors:
        build_sheet(wb, instructor,
                    courses_by_instructor.get(instructor, []),
                    office_by_instructor.get(instructor, []))

    if len(wb.sheetnames) == 0:
        wb.create_sheet("Schedule")

    wb.save(out_path)
    print(f"Wrote '{out_path}' ({len(instructors)} instructor sheet(s)).")
    return 0


if __name__ == "__main__":
    sys.exit(main())
